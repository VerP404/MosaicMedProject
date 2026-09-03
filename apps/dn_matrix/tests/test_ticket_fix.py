"""Сверка талонов для исправления услуг и выгрузка BrowserAuto."""

from __future__ import annotations

import io
import zipfile
from datetime import date

from django.test import TestCase
from openpyxl import load_workbook

from apps.dn_matrix.models import (
    DnDiagnosis,
    DnDiagnosisCategory,
    DnDiagnosisGroup,
    DnDiagnosisGroupMembership,
    DnDiagnosisSpecialty,
    DnService,
    DnServicePrice,
    DnServicePricePeriod,
    DnServiceRequirement,
    DnSpecialty,
    DnUnavailableService,
    MatrixEdition,
)
from apps.dn_matrix.services.matrix import build_matrix_pick_cache
from apps.dn_matrix.services.service_codes import ServiceCodeEquivalence
from apps.dn_matrix.services.ticket_fix import (
    DetailLine,
    build_fix_export_zip,
    extract_icd,
    is_ticket_fixable,
    parse_detail_services,
    parse_ui_date,
    reconcile_ticket_row,
    _pick_service_dates,
)


class TicketFixParseTests(TestCase):
    def test_extract_icd_from_journal_text(self):
        self.assertEqual(extract_icd("I11.9 Гипертензивная болезнь"), "I11.9")
        self.assertEqual(extract_icd("I11.9ГИПЕРТЕНЗИВНАЯ"), "I11.9")
        self.assertEqual(extract_icd("I10"), "I10")
        self.assertEqual(extract_icd("А11.9 что-то"), "A11.9")
        self.assertEqual(extract_icd("-"), "")

    def test_extract_icd_list_keeps_decimals_and_drops_prefix(self):
        from apps.dn_matrix.services.ticket_fix import drop_covered_mkb_prefixes, extract_icd_list

        self.assertEqual(
            extract_icd_list("I11.9 Гипертензивная, E11.8 СД"),
            ["I11.9", "E11.8"],
        )
        self.assertEqual(drop_covered_mkb_prefixes(["I11", "I11.9", "E11.7"]), ["I11.9", "E11.7"])

    def test_parse_ui_date(self):
        self.assertEqual(parse_ui_date("2026-04-15"), "15-04-26")
        self.assertEqual(parse_ui_date("15.04.2026"), "15-04-26")
        self.assertEqual(parse_ui_date("15-04-2026"), "15-04-26")
        self.assertEqual(parse_ui_date(date(2026, 4, 15)), "15-04-26")

    def test_parse_detail_services_csv_skips_zero_amount(self):
        raw = (
            "Талон;Код услуги;Кол-во услуг;Дата начала услуги;Дата окончания услуги;Врач в услуге;Сумма услуги\n"
            "1001;SRV1;1;01.04.2026;01.04.2026;11111;100\n"
            "1001;SRV9;1;02.04.2026;02.04.2026;11111;0\n"
            "1002;B04.047.003;2;03.04.2026;10.04.2026;22222;200\n"
        ).encode("utf-8-sig")
        by_ticket = parse_detail_services(raw, "detail_services.csv")
        self.assertEqual(len(by_ticket["1001"]), 1)
        self.assertEqual(by_ticket["1001"][0].code, "SRV1")
        self.assertEqual(by_ticket["1001"][0].start, "01-04-26")
        self.assertEqual(by_ticket["1002"][0].qty, 2)
        self.assertNotIn("SRV9", [ln.code for ln in by_ticket["1001"]])

    def test_pick_service_dates_keep_existing_and_reuse_for_new(self):
        lines = [
            DetailLine(code="SRV1", qty=1, start="01-04-26", end="01-04-26", doctor="11111"),
        ]
        equiv = ServiceCodeEquivalence()
        begin, end, doctor = _pick_service_dates(
            code="SRV1",
            title="Анализ крови",
            detail_lines=lines,
            ticket_start="01-04-26",
            ticket_end="15-04-26",
            code_equiv=equiv,
        )
        self.assertEqual((begin, end, doctor), ("01-04-26", "01-04-26", "11111"))
        begin, end, doctor = _pick_service_dates(
            code="NEW1",
            title="Новый анализ",
            detail_lines=lines,
            ticket_start="01-04-26",
            ticket_end="15-04-26",
            code_equiv=equiv,
        )
        self.assertEqual((begin, end), ("01-04-26", "01-04-26"))
        begin, end, _ = _pick_service_dates(
            code="SRV2",
            title="Диспансерный прием терапевта",
            detail_lines=lines,
            ticket_start="01-04-26",
            ticket_end="15-04-26",
            code_equiv=equiv,
        )
        self.assertEqual((begin, end), ("01-04-26", "15-04-26"))

    def test_status_2_not_fixable(self):
        self.assertFalse(
            is_ticket_fixable("2", mismatch=True, expected=[{"code": "A"}], enp="1", ds1="I10")
        )
        self.assertTrue(
            is_ticket_fixable("5", mismatch=True, expected=[{"code": "A"}], enp="1", ds1="I10")
        )
        self.assertTrue(
            is_ticket_fixable("18", mismatch=True, expected=[{"code": "A"}], enp="1", ds1="I10")
        )
        self.assertFalse(
            is_ticket_fixable("19", mismatch=True, expected=[], enp="1", ds1="I10")
        )

    def test_fmt_codes_same_order_regardless_of_input(self):
        from apps.dn_matrix.services.ticket_fix import _fmt_codes

        a = [
            {"code": "A09.05.065", "qty": 1},
            {"code": "A04.22.001", "qty": 1},
            {"code": "B01.058.002", "qty": 3},
        ]
        b = [
            {"code": "A04.22.001", "qty": 1},
            {"code": "A09.05.065", "qty": 1},
            {"code": "B01.058.002", "qty": 3},
        ]
        self.assertEqual(_fmt_codes(a), _fmt_codes(b))
        self.assertEqual(_fmt_codes(a), "A04.22.001, A09.05.065, B01.058.002×3")

    def test_norm_ticket_id_strips_excel_float(self):
        from apps.dn_matrix.services.ticket_fix import norm_ticket_id

        self.assertEqual(norm_ticket_id("22835012.0"), "22835012")
        self.assertEqual(norm_ticket_id(22835012), "22835012")
        self.assertEqual(norm_ticket_id("22835012"), "22835012")


class TicketFixSqlTests(TestCase):
    def test_goal3_sql_has_no_percent_wildcards(self):
        from apps.analytical_app.pages.head.dn.query import sql_goal3_journal_talons, sql_iszl_mkb_by_enp

        sql = sql_goal3_journal_talons(2026, date(2026, 4, 1), date(2026, 4, 30), ["Терапия"], ["1", "5", "6"])
        self.assertNotIn("%", sql)
        self.assertIn("'3'", sql)
        self.assertIn("'1'", sql)
        self.assertIn("'5'", sql)
        self.assertIn("regexp_replace", sql)
        self.assertIn("load_data_talons", sql)
        self.assertIn("load_data_complex_talons", sql)
        iszl = sql_iszl_mkb_by_enp(2026, ["1234567890123456"])
        self.assertIn("dn_app_dnline", iszl)
        self.assertIn("out_of_168n", iszl)


class TicketFixReconcileTests(TestCase):
    def setUp(self):
        self.edition = MatrixEdition.objects.create(
            code="dn_fix_test",
            title="Fix services",
            effective_from=date(2026, 4, 1),
            status=MatrixEdition.Status.ACTIVE,
            max_doctor_visits=3,
        )
        cat = DnDiagnosisCategory.objects.create(
            edition=self.edition, code="ct_4", title="БСК", sort_order=0
        )
        self.diag = DnDiagnosis.objects.create(
            edition=self.edition, category=cat, mkb_code="I10", title="Гипертония"
        )
        self.diag_sd = DnDiagnosis.objects.create(
            edition=self.edition, category=cat, mkb_code="E11.7", title="СД"
        )
        self.spec = DnSpecialty.objects.create(
            edition=self.edition, code="ther", title="Терапевт", sort_order=0
        )
        DnDiagnosisSpecialty.objects.create(diagnosis=self.diag, specialty=self.spec, kind="primary")
        DnDiagnosisSpecialty.objects.create(diagnosis=self.diag_sd, specialty=self.spec, kind="primary")

        g_bsk = DnDiagnosisGroup.objects.create(
            edition=self.edition, code="g_bsk", title="БСК", rule={"mkb_codes": ["I10"]}, sort_order=0
        )
        g_sd = DnDiagnosisGroup.objects.create(
            edition=self.edition, code="g_sd", title="СД", rule={"mkb_codes": ["E11.7"]}, sort_order=1
        )
        DnDiagnosisGroupMembership.objects.create(group=g_bsk, diagnosis=self.diag)
        DnDiagnosisGroupMembership.objects.create(group=g_sd, diagnosis=self.diag_sd)

        svc_visit = DnService.objects.create(
            edition=self.edition, code="SRV2", title="Диспансерный прием терапевта", sort_order=0
        )
        svc_lab = DnService.objects.create(
            edition=self.edition, code="SRV1", title="Анализ крови", sort_order=1
        )
        svc_sd = DnService.objects.create(
            edition=self.edition, code="A09.05.023", title="Гликированный гемоглобин", sort_order=2
        )
        DnServiceRequirement.objects.create(service=svc_visit, specialty=self.spec, diagnosis_group=g_bsk)
        DnServiceRequirement.objects.create(service=svc_lab, specialty=self.spec, diagnosis_group=g_bsk)
        DnServiceRequirement.objects.create(service=svc_sd, specialty=self.spec, diagnosis_group=g_sd)
        period = DnServicePricePeriod.objects.create(
            edition=self.edition, code="p1", title="2026", valid_from=date(2026, 4, 1), valid_to=None
        )
        for svc in (svc_visit, svc_lab, svc_sd):
            DnServicePrice.objects.create(period=period, service=svc, amount="100.00", currency="RUB")

        self.cache = build_matrix_pick_cache(self.edition)
        self.specs = list(self.edition.specialties.all())
        self.equiv = ServiceCodeEquivalence(
            [(s.id, s.code, s.title) for s in self.edition.services.all().only("id", "code", "title")]
        )

    def _raw(self, **overrides):
        row = {
            "talon": "1001",
            "enp": "1234567890123456",
            "status": "1",
            "patient": "Иванов И.И.",
            "gender": "М",
            "main_diagnosis": "I10 Гипертензия",
            "additional_diagnosis": "",
            "specialty": "Терапевт",
            "doctor": "11111 Иванов",
            "department": "Поликлиника 1",
            "treatment_start": "2026-04-01",
            "treatment_end": "2026-04-15",
            "visits": "2",
        }
        row.update(overrides)
        return row

    def test_mismatch_and_iszl_comorbid_expands_expected(self):
        item = reconcile_ticket_row(
            self._raw(),
            detail_lines=[
                DetailLine(code="SRV2", qty=2, start="01-04-26", end="15-04-26", doctor="11111"),
            ],
            iszl_mkbs={"E11.7"},
            edition=self.edition,
            cache=self.cache,
            specs=self.specs,
            equiv=self.equiv,
            price_on=date(2026, 4, 15),
        )
        codes = {e["code"] for e in item.expected}
        self.assertIn("SRV1", codes)
        self.assertIn("SRV2", codes)
        self.assertIn("A09.05.023", codes)
        self.assertIn("E11.7", item.ds2_full)
        self.assertTrue(item.fixable)
        missing_codes = {m["code"] for m in item.missing}
        self.assertIn("SRV1", missing_codes)
        self.assertIn("A09.05.023", missing_codes)

    def test_journal_ds2_keeps_full_icd_not_prefix(self):
        item = reconcile_ticket_row(
            self._raw(additional_diagnosis="I11.9 Гипертензивная болезнь сердца"),
            detail_lines=[
                DetailLine(code="SRV2", qty=2, start="01-04-26", end="15-04-26", doctor="11111"),
                DetailLine(code="SRV1", qty=1, start="03-04-26", end="03-04-26", doctor="11111"),
            ],
            iszl_mkbs={"I11", "E11.7"},
            edition=self.edition,
            cache=self.cache,
            specs=self.specs,
            equiv=self.equiv,
            price_on=date(2026, 4, 15),
        )
        self.assertEqual(item.ds2, ["I11.9"])
        self.assertNotIn("I11", item.ds2)
        self.assertNotIn("I11", item.ds2_full)
        self.assertIn("I11.9", item.ds2_full)
        self.assertIn("E11.7", item.ds2_full)
        self.assertIn("A09.05.023", {e["code"] for e in item.expected})

    def test_skip_unavailable_from_expected(self):
        DnUnavailableService.objects.create(edition=self.edition, service_code="SRV1", mkb_code="")
        cache = build_matrix_pick_cache(self.edition)
        item = reconcile_ticket_row(
            self._raw(),
            detail_lines=[
                DetailLine(code="SRV2", qty=2, start="01-04-26", end="15-04-26", doctor="11111"),
            ],
            iszl_mkbs=set(),
            edition=self.edition,
            cache=cache,
            specs=self.specs,
            equiv=self.equiv,
            price_on=date(2026, 4, 15),
        )
        codes = {e["code"] for e in item.expected}
        self.assertNotIn("SRV1", codes)
        self.assertIn("SRV2", codes)

    def test_status_2_not_exported(self):
        item = reconcile_ticket_row(
            self._raw(status="2"),
            detail_lines=[],
            iszl_mkbs=set(),
            edition=self.edition,
            cache=self.cache,
            specs=self.specs,
            equiv=self.equiv,
            price_on=date(2026, 4, 15),
        )
        self.assertTrue(item.missing)
        self.assertFalse(item.fixable)
        with self.assertRaises(ValueError):
            build_fix_export_zip([item], only_fixable=True)

    def test_zip_has_oms_ticket_and_keeps_detail_dates(self):
        item = reconcile_ticket_row(
            self._raw(),
            detail_lines=[
                DetailLine(code="SRV2", qty=2, start="01-04-26", end="15-04-26", doctor="11111"),
                DetailLine(code="SRV1", qty=1, start="03-04-26", end="03-04-26", doctor="11111"),
            ],
            iszl_mkbs={"E11.7"},
            edition=self.edition,
            cache=self.cache,
            specs=self.specs,
            equiv=self.equiv,
            price_on=date(2026, 4, 15),
        )
        # факт без A09 — mismatch, статус 1 → в пакет
        self.assertTrue(item.fixable)
        blob, stamp = build_fix_export_zip([item], only_fixable=True)
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            names = zf.namelist()
            self.assertTrue(any(n.startswith("талоны_") for n in names))
            self.assertTrue(any(n.startswith("услуги_") for n in names))
            talons_name = next(n for n in names if n.startswith("талоны_"))
            services_name = next(n for n in names if n.startswith("услуги_"))
            wb_t = load_workbook(io.BytesIO(zf.read(talons_name)))
            headers = [c.value for c in next(wb_t.active.iter_rows(min_row=1, max_row=1))]
            self.assertIn("талон_омс", headers)
            values = [c.value for c in next(wb_t.active.iter_rows(min_row=2, max_row=2))]
            row = dict(zip(headers, values))
            self.assertEqual(row["талон_омс"], "1001")
            self.assertEqual(str(row["Талон"]), "1001")
            self.assertEqual(row["Диагноз"], "I10")
            wb_s = load_workbook(io.BytesIO(zf.read(services_name)))
            svc_headers = [c.value for c in next(wb_s.active.iter_rows(min_row=1, max_row=1))]
            svc_rows = []
            for excel_row in wb_s.active.iter_rows(min_row=2, values_only=True):
                svc_rows.append(dict(zip(svc_headers, excel_row)))
            by_code = {r["Код услуги"]: r for r in svc_rows}
            self.assertTrue(svc_rows)
            self.assertTrue(all(str(r["Талон"]) == "1001" for r in svc_rows))
            self.assertEqual(by_code["SRV1"]["Дата начала"], "03-04-26")
            self.assertEqual(by_code["SRV1"]["Дата окончания"], "03-04-26")
            self.assertEqual(by_code["SRV2"]["Дата начала"], "01-04-26")
            self.assertEqual(by_code["SRV2"]["Дата окончания"], "15-04-26")
            self.assertEqual(str(by_code["SRV2"]["Кол-во"]), "2")
            self.assertEqual(by_code["A09.05.023"]["Дата начала"], "01-04-26")
        self.assertTrue(stamp)
